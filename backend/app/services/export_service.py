"""
Export search results to PDF, Excel, or Word.
"""
import io
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from docx import Document
from docx.shared import Pt, RGBColor
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from app.models import Candidate, SearchSession


RECOMMENDATION_LABELS = {
    "green": "Пригласить немедленно",
    "yellow": "Рассмотреть дополнительно",
    "red": "Низкий приоритет",
}


def export_to_excel(session: SearchSession, candidates: list[Candidate]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Кандидаты"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1E3A5F")

    headers = [
        "№", "ФИО", "Должность", "Город", "Опыт (лет)",
        "Ожидаемая зарплата", "Match Score", "Рекомендация",
        "Сильные стороны", "Риски", "Ссылка на резюме",
    ]
    ws.append(headers)
    for col, _ in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    sorted_candidates = sorted(candidates, key=lambda c: c.match_score or 0, reverse=True)

    green_fill = PatternFill("solid", fgColor="C6EFCE")
    yellow_fill = PatternFill("solid", fgColor="FFEB9C")
    red_fill = PatternFill("solid", fgColor="FFC7CE")
    fill_map = {"green": green_fill, "yellow": yellow_fill, "red": red_fill}

    for i, c in enumerate(sorted_candidates, 1):
        row = [
            i,
            c.full_name,
            c.position or "",
            c.city or "",
            c.experience_years or "",
            c.salary_expected or "",
            f"{c.match_score:.0f}%" if c.match_score else "",
            RECOMMENDATION_LABELS.get(c.recommendation or "", ""),
            "; ".join(c.pros or []),
            "; ".join(c.cons or []),
            c.hh_url or "",
        ]
        ws.append(row)
        fill = fill_map.get(c.recommendation or "")
        if fill:
            for col in range(1, len(headers) + 1):
                ws.cell(row=i + 1, column=col).fill = fill

    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 50)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def export_to_word(session: SearchSession, candidates: list[Candidate]) -> bytes:
    doc = Document()

    title = doc.add_heading(f"Виртуальный рекрутер — {session.title}", 0)
    title.runs[0].font.color.rgb = RGBColor(0x1E, 0x3A, 0x5F)

    doc.add_paragraph(f"Дата формирования: {datetime.now().strftime('%d.%m.%Y %H:%M')}")
    doc.add_paragraph(f"Всего найдено кандидатов: {len(candidates)}")
    doc.add_paragraph()

    sorted_candidates = sorted(candidates, key=lambda c: c.match_score or 0, reverse=True)

    for i, c in enumerate(sorted_candidates[:20], 1):
        doc.add_heading(f"{i}. {c.full_name} — {c.match_score:.0f}%" if c.match_score else f"{i}. {c.full_name}", 2)

        table = doc.add_table(rows=1, cols=2)
        table.style = "Table Grid"
        info = [
            ("Должность", c.position or "—"),
            ("Город", c.city or "—"),
            ("Опыт", f"{c.experience_years} лет" if c.experience_years else "—"),
            ("Зарплата", f"{c.salary_expected:,} ₸" if c.salary_expected else "—"),
            ("Рекомендация", RECOMMENDATION_LABELS.get(c.recommendation or "", "—")),
        ]
        hdr = table.rows[0].cells
        hdr[0].text = "Параметр"
        hdr[1].text = "Значение"
        for label, value in info:
            row = table.add_row().cells
            row[0].text = label
            row[1].text = value

        if c.pros:
            doc.add_paragraph("Почему подходит:", style="Intense Quote")
            for pro in c.pros:
                doc.add_paragraph(f"✓ {pro}", style="List Bullet")

        if c.cons:
            doc.add_paragraph("Риски:", style="Intense Quote")
            for con in c.cons:
                doc.add_paragraph(f"⚠ {con}", style="List Bullet")

        if c.ai_summary:
            doc.add_paragraph(c.ai_summary)

        if c.hh_url:
            doc.add_paragraph(f"Резюме: {c.hh_url}")

        doc.add_paragraph()

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def export_to_pdf(session: SearchSession, candidates: list[Candidate]) -> bytes:
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, rightMargin=30, leftMargin=30, topMargin=30, bottomMargin=30)
    styles = getSampleStyleSheet()

    story = []
    title_style = ParagraphStyle("title", parent=styles["Title"], fontSize=16, textColor=colors.HexColor("#1E3A5F"))
    story.append(Paragraph(f"Виртуальный рекрутер — {session.title}", title_style))
    story.append(Spacer(1, 12))
    story.append(Paragraph(f"Сформировано: {datetime.now().strftime('%d.%m.%Y %H:%M')}", styles["Normal"]))
    story.append(Spacer(1, 20))

    sorted_candidates = sorted(candidates, key=lambda c: c.match_score or 0, reverse=True)

    table_data = [["№", "ФИО", "Должность", "Опыт", "Score", "Рекомендация"]]
    for i, c in enumerate(sorted_candidates[:20], 1):
        table_data.append([
            str(i),
            c.full_name[:30],
            (c.position or "")[:25],
            f"{c.experience_years} л." if c.experience_years else "—",
            f"{c.match_score:.0f}%" if c.match_score else "—",
            RECOMMENDATION_LABELS.get(c.recommendation or "", "—")[:20],
        ])

    t = Table(table_data, colWidths=[25, 120, 110, 45, 45, 115])
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1E3A5F")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F5F5F5")]),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
    ]))
    story.append(t)

    doc.build(story)
    return buffer.getvalue()
